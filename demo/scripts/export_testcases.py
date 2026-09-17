"""Export Excel IBox cases and overlay gold labels grounded on the 8 PDFs."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = Path(r"E:\ibox_qarr\files") / "制度查询智能体测试用例（出差方面）及测试结果.xlsx"
OUT = ROOT / "data" / "testcases.json"

GOLD = {
    1: {
        "bucket": "库内无据类",
        "gold_doc_ids": [],
        "gold_clause_ids": [],
        "related_clause_ids": ["3-1-14:第六条"],
        "gold_must_include": ["未在制度库中找到", "综合部"],
        "gold_must_not": ["付款管理规定"],
        "unanswerable": True,
        "notes": "原 IBox 引用的《付款管理规定》不在本 DEMO 8 份文件中。报销管理办法仅写预算员检查预算，未规定“无法提交报销流程”的处理步骤。正确表现是承认无据并建议咨询部门预算管理员/综合部。",
    },
    2: {
        "bucket": "跨文件类",
        "gold_doc_ids": ["3-1-13", "3-1-14"],
        "gold_clause_ids": ["3-1-13:第八条", "3-1-13:第十三条", "3-1-13:第十六条", "3-1-14:第十条"],
        "gold_must_include": ["公共交通", "80", "过路费"],
        "gold_must_not": ["付款管理规定"],
        "unanswerable": False,
        "notes": "2022 年版本市往返机场出租车时段为 9:30-17:30 转地铁，不是评测表 8 点口径。库内未写网约车行程单。公派车辆见第十三条。",
    },
    3: {
        "bucket": "跨文件类",
        "gold_doc_ids": ["3-1-13", "3-1-14"],
        "gold_clause_ids": ["3-1-13:第十七条", "3-1-14:第十条"],
        "gold_must_include": ["八项规定", "个人承担"],
        "gold_must_not": [],
        "unanswerable": False,
        "notes": "出差管理规定 1.1 修订删除了“主办方指定酒店”条款，但《报销管理办法》第十条仍保留书面文件例外。超标原则见出差第十七条（三）：超出部分个人承担。",
    },
    4: {
        "bucket": "精确数字与条款类",
        "gold_doc_ids": ["3-1-13", "3-1-14"],
        "gold_clause_ids": ["3-1-13:第十三条", "3-1-14:第十条"],
        "gold_must_include": ["不得报销", "过路费"],
        "gold_must_not": [],
        "unanswerable": False,
        "notes": "公司提供交通工具期间不得报销车费或市内交通费，过路费/通行费除外。",
    },
    5: {
        "bucket": "精确数字与条款类",
        "gold_doc_ids": ["3-1-13", "3-1-14", "3-1-21"],
        "gold_clause_ids": ["3-1-13:第五条", "3-1-14:第十条", "3-1-21:第八条"],
        "gold_must_include": ["提前", "出差申请"],
        "gold_must_not": ["公派车辆"],
        "unanswerable": False,
        "notes": "报销管理办法第十条写明出差前至少提前 2 天填写申请。综合部“订机票提前 1 天”库内无对应原文。禁止串题答成公派车辆交通费。",
    },
    6: {
        "bucket": "模糊场景类",
        "gold_doc_ids": ["3-1-13", "3-1-21"],
        "gold_clause_ids": ["3-1-13:第五条", "3-1-21:第八条"],
        "gold_must_include": ["审批", "出差申请"],
        "gold_must_not": ["一个工作日内完成补办"],
        "unanswerable": False,
        "notes": "出差须事先审批，商旅预定无前置审批不予报销。日常业务支出第十三条“紧急用车一个工作日内补办”是出租车/用车，不得套用到出差。",
    },
    7: {
        "bucket": "精确数字与条款类",
        "gold_doc_ids": ["3-1-13"],
        "gold_clause_ids": ["3-1-13:第八条", "3-1-13:第九条", "3-1-13:第十二条", "3-1-13:附件1"],
        "gold_must_include": ["北京", "广州", "深圳", "120"],
        "gold_must_not": [],
        "unanswerable": False,
        "has_table": True,
        "notes": "必须读第九条住宿表和附件1 城市清单原文，禁止编造二线城市名单。津贴时段以 17:30 为准。",
    },
    8: {
        "bucket": "精确数字与条款类",
        "gold_doc_ids": ["3-1-13", "3-1-14"],
        "gold_clause_ids": ["3-1-13:第九条", "3-1-13:第十七条", "3-1-14:第十条"],
        "gold_must_include": ["个人承担"],
        "gold_must_not": [],
        "unanswerable": False,
        "notes": "超标原则上个人承担；旺季上浮 20%、重大活动可工作请示；报销管理办法允许特殊情况事先工作请示。",
    },
    9: {
        "bucket": "术语口语类",
        "gold_doc_ids": ["3-1-13"],
        "gold_clause_ids": ["3-1-13:第九条"],
        "gold_must_include": ["工作请示", "证明"],
        "gold_must_not": [],
        "unanswerable": False,
        "notes": "当地重大活动住宿普涨：可申请酌情上浮，需相关证明，以工作请示审批通过为准。",
    },
    10: {
        "bucket": "术语口语类",
        "gold_doc_ids": ["3-1-13"],
        "gold_clause_ids": ["3-1-13:第九条", "3-1-13:附件2"],
        "gold_must_include": ["20%", "旺季"],
        "gold_must_not": [],
        "unanswerable": False,
        "has_table": True,
        "notes": "旅游旺季对应附件2 季节性城市，限额最高上浮 20%。",
    },
    11: {
        "bucket": "术语口语类",
        "gold_doc_ids": ["3-1-13", "3-1-21"],
        "gold_clause_ids": ["3-1-13:第七条", "3-1-21:第四条", "3-1-21:第六条"],
        "gold_must_include": ["商旅", "钉钉"],
        "gold_must_not": [],
        "unanswerable": False,
        "notes": "统一商旅软件，钉钉关联进入。",
    },
    12: {
        "bucket": "模糊场景类",
        "gold_doc_ids": ["3-1-21"],
        "gold_clause_ids": ["3-1-21:第十条"],
        "gold_must_include": ["综合管理部", "另行预定"],
        "gold_must_not": [],
        "unanswerable": False,
        "notes": "所在城市商旅软件无符合规定酒店：报告综合管理部，经批准后可另行预定。",
    },
    13: {
        "bucket": "精确数字与条款类",
        "gold_doc_ids": ["3-1-21"],
        "gold_clause_ids": ["3-1-21:第九条"],
        "gold_must_include": ["工作请示", "不予报销"],
        "gold_must_not": [],
        "unanswerable": False,
        "notes": "必须走指定商旅软件；无特殊原因非经工作请示不得另行购票，否则不予报销。",
    },
    14: {
        "bucket": "跨文件类",
        "gold_doc_ids": ["3-1-13", "3-1-14"],
        "gold_clause_ids": ["3-1-13:第十一条", "3-1-13:第十七条", "3-1-14:第十条", "3-1-14:第十三条", "3-1-14:第十四条"],
        "gold_must_include": ["差旅费报销", "业务招待"],
        "gold_must_not": [],
        "unanswerable": False,
        "notes": "差旅走差旅费报销并附加出差申请；出差期间业务招待单独走业务招待；公司组织会议走会务费。",
    },
    15: {
        "bucket": "精确数字与条款类",
        "gold_doc_ids": ["3-1-13", "3-1-14"],
        "gold_clause_ids": ["3-1-13:第六条", "3-1-14:第十一条"],
        "gold_must_include": ["1元", "公里"],
        "gold_must_not": [],
        "unanswerable": False,
        "policy_vs_hr": "原文允许特殊情况私车公用并按 1 元/公里补贴；综合部口径为目前禁止私车公用。DEMO 验收看原文准确，可提示现行执行请咨询综合部。",
        "notes": "公司不提倡私车因公离沪；特殊情况事先申请。油补 1 元/公里，高速/停车按实报。",
    },
    16: {
        "bucket": "精确数字与条款类",
        "gold_doc_ids": ["3-1-14"],
        "gold_clause_ids": ["3-1-14:第二十一条", "3-1-14:第二十二条"],
        "gold_must_include": ["3个月"],
        "gold_must_not": [],
        "unanswerable": False,
        "notes": "同一会计年度及跨年费用原则上 3 个月内报销，逾期 3 个月以上原则上不予报销。",
    },
    17: {
        "bucket": "精确数字与条款类",
        "gold_doc_ids": ["3-1-13"],
        "gold_clause_ids": ["3-1-13:第十二条"],
        "gold_must_include": ["120", "17:30"],
        "gold_must_not": ["19:30"],
        "unanswerable": False,
        "has_table": True,
        "notes": "2022 年版津贴表以 17:30 为界，不是 2025 年版的 19:30。宾馆 120 元/天。",
    },
}


def parse_ibox_score(text: str) -> str | None:
    if not text:
        return None
    m = re.search(r"准确度[：:]\s*(\d+%)", text)
    return m.group(1) if m else None


def parse_ibox_verdict(text: str) -> str:
    if not text:
        return ""
    m = re.search(r"结论[：:]\s*([^\n，,]+)", text)
    return m.group(1).strip() if m else text[:40]


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    wb = load_workbook(XLSX, data_only=True)
    ws = wb.active
    cases = []
    for row in range(3, ws.max_row + 1):
        seq = ws[f"A{row}"].value
        if not seq:
            continue
        seq = int(seq)
        gold = GOLD[seq]
        ibox = str(ws[f"D{row}"].value or "")
        eval_txt = str(ws[f"F{row}"].value or "")
        case = {
            "id": f"tc-{seq:02d}",
            "seq": seq,
            "category": str(ws[f"B{row}"].value or "").replace("\n", ""),
            "question": str(ws[f"C{row}"].value or "").strip(),
            "ibox_answer": ibox.strip(),
            "ibox_files": str(ws[f"E{row}"].value or "").strip(),
            "ibox_eval": eval_txt.strip(),
            "ibox_score": parse_ibox_score(eval_txt),
            "ibox_verdict": parse_ibox_verdict(eval_txt),
            "hr_reply": str(ws[f"G{row}"].value or "").strip(),
            **gold,
        }
        cases.append(case)

    payload = {
        "source": XLSX.name,
        "corpus": "files/出差相关制度 8 PDFs (2022/2018)",
        "labeling_principle": "gold 以 8 份库内原文为准；综合部口径单独记录在 hr_reply / policy_vs_hr",
        "cases": cases,
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {len(cases)} cases -> {OUT}")
    for c in cases:
        print(f"{c['id']} [{c['bucket']}] gold={c['gold_clause_ids'] or 'UNANSWERABLE'} score={c['ibox_score']}")


if __name__ == "__main__":
    main()
